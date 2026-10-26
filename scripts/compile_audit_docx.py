"""Compile TRACEHUNT_COMPREHENSIVE_AUDIT_REPORT.md into a styled Word (.docx) report."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent.parent
MD_PATH = BASE_DIR / "reports" / "TRACEHUNT_COMPREHENSIVE_AUDIT_REPORT.md"
DOCX_PATH = BASE_DIR / "reports" / "TRACEHUNT_COMPREHENSIVE_AUDIT_REPORT.docx"


def set_cell_background(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def main():
    print(f"[+] Reading {MD_PATH}...")
    lines = MD_PATH.read_text(encoding="utf-8").split("\n")

    doc = Document()

    # 1 inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    in_table = False
    table_rows = []
    in_code = False
    code_lines = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        num_cols = max(len(r) for r in table_rows)
        tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True

        for r_idx, row_data in enumerate(table_rows):
            is_header = (r_idx == 0)
            for c_idx, val in enumerate(row_data):
                if c_idx < num_cols:
                    cell = tbl.cell(r_idx, c_idx)
                    cell.text = val
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    if is_header:
                        set_cell_background(cell, "1E3A8A")
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.font.size = Pt(9.5)
                    else:
                        if r_idx % 2 == 1:
                            set_cell_background(cell, "F8FAFC")
                        for run in p.runs:
                            run.font.size = Pt(9)
        doc.add_paragraph()  # spacing
        in_table = False
        table_rows = []

    def flush_code():
        nonlocal in_code, code_lines
        if not code_lines:
            in_code = False
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run("\n".join(code_lines))
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
        in_code = False
        code_lines = []

    for line in lines:
        stripped = line.strip()

        # Code block
        if stripped.startswith("```"):
            if in_code:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code = True
                code_lines = []
            continue

        if in_code:
            code_lines.append(line)
            continue

        # Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [re.sub(r"\*\*(.+?)\*\*", r"\1", c.strip()) for c in stripped[1:-1].split("|")]
            # Skip separator line
            if all(re.match(r"^:?-+:?$", c) for c in cells):
                continue
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(cells)
            continue
        else:
            if in_table:
                flush_table()

        if not stripped:
            continue

        # Headings
        if stripped.startswith("# "):
            p = doc.add_heading(stripped[2:], level=1)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif stripped.startswith("## "):
            p = doc.add_heading(stripped[3:], level=2)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
        elif stripped.startswith("### "):
            p = doc.add_heading(stripped[4:], level=3)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
        elif stripped.startswith("#### "):
            p = doc.add_heading(stripped[5:], level=4)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        elif stripped == "---":
            pass
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(stripped[2:], style="List Bullet")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
        elif re.match(r"^\d+\.\s", stripped):
            item_text = re.sub(r"^\d+\.\s", "", stripped)
            p = doc.add_paragraph(item_text, style="List Number")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
        else:
            clean_text = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1", stripped)
            p = doc.add_paragraph(clean_text)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)

    if in_table:
        flush_table()
    if in_code:
        flush_code()

    print(f"[+] Saving Word document to {DOCX_PATH}...")
    doc.save(str(DOCX_PATH))
    print(f"[+] SUCCESS! Generated Word audit report: {DOCX_PATH} ({DOCX_PATH.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
