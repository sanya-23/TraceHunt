"""Microsoft Word (.docx) Investigation Report Exporter.

Renders high-fidelity, executive-grade Word investigation briefings using
python-docx with structured tables, severity highlights, and audit trails.
"""

from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from tracehunt.config import settings
from tracehunt.logging_config import get_logger
from tracehunt.reporting.schema import ReportData

logger = get_logger("reporting.docx")


def _set_cell_background(cell, hex_color: str) -> None:
    """Set the background color of a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150) -> None:
    """Set inner margins (padding) for a table cell in dxa units."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def export_docx_report(report_data: ReportData, filename: str | None = None) -> Path:
    """Export ReportData to a professionally styled Microsoft Word document (.docx)."""
    settings.REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    report_id = report_data.metadata.report_id
    safe_filename = filename if filename else f"{report_id}.docx"
    if not safe_filename.endswith(".docx"):
        safe_filename += ".docx"

    out_path = (settings.REPORTS_DIR / safe_filename).resolve()

    doc = Document()

    # Page Margins: 1 inch on all sides
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 1. Document Header / Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_pre = title_p.add_run("TRACEHUNT ENDPOINT THREAT HUNTING PLATFORM\n")
    run_pre.font.name = "Calibri"
    run_pre.font.size = Pt(10)
    run_pre.font.bold = True
    run_pre.font.color.rgb = RGBColor(30, 64, 175)  # Navy Blue

    run_title = title_p.add_run(report_data.metadata.title or "Executive Threat Investigation Briefing")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(16)
    sub_run = subtitle_p.add_run(
        f"Classification: TLP:AMBER | Confidential Security Operations Investigation\n"
        f"Report ID: {report_id}  |  Generated: {report_data.metadata.generated_at}"
    )
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(9.5)
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    # 2. Executive Scope & Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_items = [
        ("Target Host", report_data.metadata.host or "Enterprise Multi-Host"),
        ("Investigation ID", report_data.metadata.investigation_id or "Ad-hoc Telemetry Review"),
        ("Lead Analyst", report_data.metadata.analyst_name or "analyst_ui"),
        ("Time Window", f"{report_data.metadata.time_range_start or 'Initial Ingestion'} to {report_data.metadata.time_range_end or 'Present'}"),
    ]

    for row_idx, (k, v) in enumerate(meta_items):
        cell_k = meta_table.cell(row_idx, 0)
        cell_v = meta_table.cell(row_idx, 1)
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        _set_cell_background(cell_k, "F1F5F9")
        _set_cell_background(cell_v, "FFFFFF")
        _set_cell_margins(cell_k, top=60, bottom=60, left=100, right=100)
        _set_cell_margins(cell_v, top=60, bottom=60, left=100, right=100)

        p_k = cell_k.paragraphs[0]
        p_k.paragraph_format.space_after = Pt(0)
        r_k = p_k.add_run(k)
        r_k.font.bold = True
        r_k.font.size = Pt(9.5)

        p_v = cell_v.paragraphs[0]
        p_v.paragraph_format.space_after = Pt(0)
        r_v = p_v.add_run(str(v))
        r_v.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 3. Investigation Summary (If Case Attached)
    if report_data.case:
        h2 = doc.add_heading("1. Investigation Case Summary", level=1)
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(4)

        p_case = doc.add_paragraph()
        p_case.paragraph_format.space_after = Pt(8)
        r_title = p_case.add_run(f"Case {report_data.case.id}: {report_data.case.title} [{report_data.case.status}]\n")
        r_title.font.bold = True
        r_title.font.size = Pt(11)

        r_sum = p_case.add_run(report_data.case.summary or "No initial hypothesis provided.")
        r_sum.font.size = Pt(10)
        r_sum.font.italic = True

    # 4. Telemetry & Threat Metrics
    h_metrics = doc.add_heading("2. Operational Telemetry & Finding Metrics", level=1)
    h_metrics.paragraph_format.space_before = Pt(14)
    h_metrics.paragraph_format.space_after = Pt(6)

    m_table = doc.add_table(rows=2, cols=4)
    m_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    metrics_headers = ["Total Events", "Total Findings", "Critical / High", "Active Detection Rules"]
    crit_high = report_data.metrics.severity_counts.get("CRITICAL", 0) + report_data.metrics.severity_counts.get("HIGH", 0)
    metrics_vals = [
        str(report_data.metrics.total_events),
        str(report_data.metrics.total_findings),
        str(crit_high),
        str(len(report_data.metrics.detection_rules_triggered)),
    ]

    for col_idx, (hdr, val) in enumerate(zip(metrics_headers, metrics_vals)):
        cell_h = m_table.cell(0, col_idx)
        cell_v = m_table.cell(1, col_idx)
        cell_h.width = Inches(1.6)
        cell_v.width = Inches(1.6)
        _set_cell_background(cell_h, "1E3A8A")  # Dark Blue
        _set_cell_background(cell_v, "F8FAFC")
        _set_cell_margins(cell_h, top=80, bottom=80, left=100, right=100)
        _set_cell_margins(cell_v, top=100, bottom=100, left=100, right=100)

        p_h = cell_h.paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h.paragraph_format.space_after = Pt(0)
        rh = p_h.add_run(hdr)
        rh.font.bold = True
        rh.font.color.rgb = RGBColor(255, 255, 255)
        rh.font.size = Pt(9)

        p_v = cell_v.paragraphs[0]
        p_v.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_v.paragraph_format.space_after = Pt(0)
        rv = p_v.add_run(val)
        rv.font.bold = True
        rv.font.size = Pt(16)
        rv.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 5. Detected Behavioral Findings Section
    h_find = doc.add_heading(f"3. Explainable Threat Findings ({len(report_data.findings)} Total)", level=1)
    h_find.paragraph_format.space_before = Pt(14)
    h_find.paragraph_format.space_after = Pt(6)

    if not report_data.findings:
        p_none = doc.add_paragraph("No behavioral detection findings were registered for this scope.")
        p_none.runs[0].font.italic = True
    else:
        for f_idx, f in enumerate(report_data.findings, 1):
            p_f = doc.add_paragraph()
            p_f.paragraph_format.space_before = Pt(8)
            p_f.paragraph_format.space_after = Pt(2)

            sev_colors = {
                "CRITICAL": RGBColor(220, 38, 38),
                "HIGH": RGBColor(234, 88, 12),
                "MEDIUM": RGBColor(202, 138, 4),
                "LOW": RGBColor(37, 99, 235),
            }
            c_rgb = sev_colors.get(f.severity.upper(), RGBColor(100, 116, 139))

            r_badge = p_f.add_run(f"[{f.severity}] ")
            r_badge.font.bold = True
            r_badge.font.color.rgb = c_rgb

            r_hdr = p_f.add_run(f"Finding #{f.id}: {f.rule_id} on {f.host}")
            r_hdr.font.bold = True
            r_hdr.font.size = Pt(11)

            # Detail table for finding
            f_table = doc.add_table(rows=4, cols=2)
            f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            f_rows = [
                ("Detection Rationale", f.rationale),
                ("Process & Command", f"{f.process_image or 'N/A'}\n{f.command_line or 'N/A'}"),
                ("Matched Indicators", str(f.matched_fields) if f.matched_fields else "N/A"),
                ("Caveats / Guidance", f.caveats or "Verified heuristic rule trigger."),
            ]
            for r_i, (k, val) in enumerate(f_rows):
                ck = f_table.cell(r_i, 0)
                cv = f_table.cell(r_i, 1)
                ck.width = Inches(2.0)
                cv.width = Inches(4.5)
                _set_cell_background(ck, "F8FAFC")
                _set_cell_margins(ck, top=40, bottom=40, left=80, right=80)
                _set_cell_margins(cv, top=40, bottom=40, left=80, right=80)

                pk = ck.paragraphs[0]
                pk.paragraph_format.space_after = Pt(0)
                rk = pk.add_run(k)
                rk.font.bold = True
                rk.font.size = Pt(9)

                pv = cv.paragraphs[0]
                pv.paragraph_format.space_after = Pt(0)
                rv = pv.add_run(val)
                rv.font.size = Pt(9)

            doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 6. Chronological Host Timeline Summary
    h_time = doc.add_heading(f"4. Chronological Telemetry Timeline (Sample: First {min(len(report_data.timeline), 30)} Events)", level=1)
    h_time.paragraph_format.space_before = Pt(14)
    h_time.paragraph_format.space_after = Pt(6)

    if report_data.timeline:
        t_table = doc.add_table(rows=1, cols=4)
        t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_hdrs = ["Event Time (UTC)", "Type", "Image / User", "Command Line / Details"]
        t_widths = [Inches(1.8), Inches(1.3), Inches(1.6), Inches(1.8)]

        for c_i, (h_text, w) in enumerate(zip(t_hdrs, t_widths)):
            c = t_table.cell(0, c_i)
            c.width = w
            _set_cell_background(c, "334155")  # Slate Gray
            _set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h_text)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.5)

        for item in report_data.timeline[:30]:
            row = t_table.add_row()
            t_row_vals = [
                item.event_time[:19].replace("T", " ") if item.event_time else "N/A",
                item.event_type,
                f"{Path(item.image).name if item.image else 'N/A'}\n({item.user or 'N/A'})",
                (item.command_line[:120] + "...") if item.command_line and len(item.command_line) > 120 else (item.command_line or "N/A"),
            ]
            for c_i, (val, w) in enumerate(zip(t_row_vals, t_widths)):
                c = row.cells[c_i]
                c.width = w
                _set_cell_margins(c, top=40, bottom=40, left=80, right=80)
                if item.has_detection:
                    _set_cell_background(c, "FEF2F2")  # Light Red highlight
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(val)
                r.font.size = Pt(8)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 7. Case Audit Trail (If Case Attached)
    if report_data.case and report_data.case.audit_trail:
        h_audit = doc.add_heading("5. Chain of Custody & Analyst Audit Trail", level=1)
        h_audit.paragraph_format.space_before = Pt(14)
        h_audit.paragraph_format.space_after = Pt(6)

        a_table = doc.add_table(rows=1, cols=4)
        a_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        a_hdrs = ["Timestamp (UTC)", "Action", "Reviewer", "Comment / Disposition Details"]
        a_widths = [Inches(1.8), Inches(1.3), Inches(1.2), Inches(2.2)]

        for c_i, (h_text, w) in enumerate(zip(a_hdrs, a_widths)):
            c = a_table.cell(0, c_i)
            c.width = w
            _set_cell_background(c, "1E293B")
            _set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h_text)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.5)

        for audit in report_data.case.audit_trail:
            row = a_table.add_row()
            a_vals = [
                (audit.get("created_at") or "")[:19].replace("T", " "),
                audit.get("action", "ACTION"),
                audit.get("reviewer", "N/A"),
                audit.get("comment", "") or f"Status: {audit.get('previous_status')} -> {audit.get('new_status')}",
            ]
            for c_i, (val, w) in enumerate(zip(a_vals, a_widths)):
                c = row.cells[c_i]
                c.width = w
                _set_cell_margins(c, top=40, bottom=40, left=80, right=80)
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(val)
                r.font.size = Pt(8)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 8. Forensic Disclosures & Limitations
    h_lim = doc.add_heading("6. Forensic Disclosures & Operational Limitations", level=1)
    h_lim.paragraph_format.space_before = Pt(14)
    h_lim.paragraph_format.space_after = Pt(6)

    for i, lim in enumerate(report_data.limitations, 1):
        p_l = doc.add_paragraph(lim, style="List Number")
        p_l.paragraph_format.space_after = Pt(4)
        p_l.runs[0].font.size = Pt(9)
        p_l.runs[0].font.color.rgb = RGBColor(71, 85, 105)

    doc.save(str(out_path))
    logger.info(f"Rendered executive Word report (.docx) to {out_path} ({out_path.stat().st_size} bytes)")
    return out_path
